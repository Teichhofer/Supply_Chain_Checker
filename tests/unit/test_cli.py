"""Tests for CLI entry points."""

from __future__ import annotations

import json
import os
import re
import stat
from pathlib import Path

import pytest

from supply_chain_checker import cli
from supply_chain_checker.models import ExtractedProduct
from openpyxl import Workbook, load_workbook

from supply_chain_checker.services.csv_service import StorageIOError
from supply_chain_checker.services.llm.base import LlmConfigurationError
from supply_chain_checker.services.pdf_reader import PdfProcessingError
from supply_chain_checker.services.status_service import StatusTrackingError

_MINIMAL_CONFIG = """
paths:
  input_dir: data/input
  output_dir: data/output
  state_dir: data/state
  logs_dir: logs
logging:
  level: INFO
  file_name: app.log
llm:
  provider: openai
  model: gpt-4.1-mini
  timeout_seconds: 30
  max_retries: 2
  temperature: 0.0
prompts:
  extraction: "Extract from {document_name}: {document_text}"
  assessment: "Assess {product_name} from {supplier}"
parameters:
  use_ocr_fallback: true
  max_products_per_document: 100
  max_assessment_reason_words: 100
  on_corrupt_status_file: abort
"""



def _write_extraction_xlsx(path: Path, rows: list[list[str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "run_id",
        "document_name",
        "product_name",
        "quantity",
        "supplier",
        "manufacturer",
        "article_number",
        "extraction_status",
        "extraction_hint",
    ])
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _read_xlsx_rows(path: Path) -> list[tuple[object, ...]]:
    return list(load_workbook(path, read_only=True, data_only=True).active.iter_rows(values_only=True))

_FALLBACK_STATUS_CONFIG = """
paths:
  input_dir: data/input
  output_dir: data/output
  state_dir: data/state
  logs_dir: logs
logging:
  level: INFO
  file_name: app.log
llm:
  provider: openai
  model: gpt-4.1-mini
  timeout_seconds: 30
  max_retries: 2
  temperature: 0.0
prompts:
  extraction: "Extract from {document_name}: {document_text}"
  assessment: "Assess {product_name} from {supplier}"
parameters:
  use_ocr_fallback: true
  max_products_per_document: 100
  max_assessment_reason_words: 100
  on_corrupt_status_file: fallback_empty
"""


def test_build_parser_supports_extract_assess_and_run() -> None:
    parser = cli._build_parser()

    extract_args = parser.parse_args(["extract", "--config", "config/config.yaml"])
    assess_args = parser.parse_args(["assess", "--config", "config/config.yaml"])
    run_args = parser.parse_args(["run", "--config", "config/config.yaml"])
    extract_default_args = parser.parse_args(["extract"])
    assess_default_args = parser.parse_args(["assess"])
    run_default_args = parser.parse_args(["run"])

    assert extract_args.command == "extract"
    assert assess_args.command == "assess"
    assert run_args.command == "run"
    assert extract_default_args.config == str(cli.DEFAULT_CONFIG_PATH)
    assert assess_default_args.config == str(cli.DEFAULT_CONFIG_PATH)
    assert run_default_args.config == str(cli.DEFAULT_CONFIG_PATH)



def test_build_parser_supports_clear() -> None:
    parser = cli._build_parser()

    clear_args = parser.parse_args(["clear", "--config", "config/config.yaml"])
    clear_default_args = parser.parse_args(["clear"])

    assert clear_args.command == "clear"
    assert clear_default_args.config == str(cli.DEFAULT_CONFIG_PATH)


def test_main_uses_default_config_path_when_not_provided(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_dir = tmp_path / "config"
    config_dir.mkdir(parents=True, exist_ok=True)
    (config_dir / "config.yaml").write_text(_MINIMAL_CONFIG, encoding="utf-8")

    monkeypatch.setattr("sys.argv", ["supply-chain-checker", "extract"])

    assert cli.main() == 0


def test_main_ensures_layout_and_returns_success(monkeypatch, tmp_path: Path, capsys) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "extract", "--config", str(config_file)]
    )

    exit_code = cli.main()

    assert exit_code == 0
    assert (tmp_path / "data" / "input").exists()
    assert (tmp_path / "tests" / "unit").exists()
    assert (tmp_path / "logs" / "app.log").exists()

    stdout = capsys.readouterr().out
    assert "scaffold is ready" in stdout
    assert "run_id='" in stdout

    run_id_match = re.search(r"run_id='([a-f0-9]{12})'", stdout)
    assert run_id_match is not None
    run_id = run_id_match.group(1)

    output_files = list((tmp_path / "data" / "output").glob("extraction_*.xlsx"))
    assert len(output_files) == 1

    rows = _read_xlsx_rows(output_files[0])
    assert rows[0][0] == "run_id"
    assert rows[1][0] == run_id



def test_main_clear_deletes_logs_output_and_state_directories(
    monkeypatch, tmp_path: Path
) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    logs_dir = tmp_path / "logs"
    logs_dir.mkdir(parents=True, exist_ok=True)
    (logs_dir / "app.log").write_text("log", encoding="utf-8")

    output_dir = tmp_path / "data" / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "artifact.xlsx").write_text("csv", encoding="utf-8")

    state_dir = tmp_path / "data" / "state"
    state_dir.mkdir(parents=True, exist_ok=True)
    (state_dir / "processed_files.json").write_text("{}", encoding="utf-8")

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "clear", "--config", str(config_file)]
    )

    assert cli.main() == 0
    assert logs_dir.exists()
    assert output_dir.exists()
    assert state_dir.exists()
    assert list(logs_dir.iterdir()) == []
    assert list(output_dir.iterdir()) == []
    assert list(state_dir.iterdir()) == []


def test_main_clear_deletes_nested_state_files(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    state_nested_dir = tmp_path / "data" / "state" / "nested"
    state_nested_dir.mkdir(parents=True, exist_ok=True)
    (state_nested_dir / "a.json").write_text("{}", encoding="utf-8")
    (tmp_path / "data" / "state" / ".hidden").write_text("x", encoding="utf-8")

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "clear", "--config", str(config_file)]
    )

    assert cli.main() == 0
    state_dir = tmp_path / "data" / "state"
    assert state_dir.exists()
    assert list(state_dir.iterdir()) == []


def test_main_clear_deletes_legacy_processed_files_file(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")
    (tmp_path / "processed_files.json").write_text("{}", encoding="utf-8")

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "clear", "--config", str(config_file)]
    )

    assert cli.main() == 0
    assert not (tmp_path / "processed_files.json").exists()


def test_handle_remove_readonly_does_not_recurse_to_rmtree(monkeypatch, tmp_path: Path) -> None:
    nested_dir = tmp_path / "state"
    nested_dir.mkdir(parents=True, exist_ok=True)
    (nested_dir / "artifact.txt").write_text("x", encoding="utf-8")

    def _unexpected_rmtree(*_args: object, **_kwargs: object) -> None:
        raise AssertionError("shutil.rmtree must not be called from onerror handler")

    monkeypatch.setattr(cli.shutil, "rmtree", _unexpected_rmtree)

    cli._handle_remove_readonly(os.rmdir, str(nested_dir), PermissionError("permission denied"))


def test_handle_remove_readonly_unlinks_readonly_file(tmp_path: Path) -> None:
    target_file = tmp_path / "readonly.log"
    target_file.write_text("log", encoding="utf-8")
    target_file.chmod(stat.S_IREAD)

    cli._handle_remove_readonly(os.unlink, str(target_file), PermissionError("permission denied"))

    assert not target_file.exists()


def test_main_loads_openai_key_from_sibling_secrets_env(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)

    config_dir = tmp_path / "config"
    config_dir.mkdir(parents=True, exist_ok=True)
    (config_dir / "config.yaml").write_text(_MINIMAL_CONFIG, encoding="utf-8")
    (config_dir / "secrets.env").write_text(
        "export OPENAI_API_KEY='sk-test-from-secrets'\n",
        encoding="utf-8",
    )

    monkeypatch.setattr(
        "sys.argv",
        ["supply-chain-checker", "extract", "--config", str(config_dir / "config.yaml")],
    )

    assert cli.main() == 0
    assert os.getenv("OPENAI_API_KEY") == "sk-test-from-secrets"
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)


def test_main_prefers_openai_key_from_sibling_secrets_env(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("OPENAI_API_KEY", "sk-existing")

    config_dir = tmp_path / "config"
    config_dir.mkdir(parents=True, exist_ok=True)
    (config_dir / "config.yaml").write_text(_MINIMAL_CONFIG, encoding="utf-8")
    (config_dir / "secrets.env").write_text(
        "OPENAI_API_KEY=sk-from-secrets\n",
        encoding="utf-8",
    )

    monkeypatch.setattr(
        "sys.argv",
        ["supply-chain-checker", "extract", "--config", str(config_dir / "config.yaml")],
    )

    assert cli.main() == 0
    assert os.getenv("OPENAI_API_KEY") == "sk-from-secrets"


def test_main_logs_run_finished_even_when_command_raises(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "extract", "--config", str(config_file)]
    )

    def _raise_runtime_error(*_args: object, **_kwargs: object) -> None:
        raise RuntimeError

    monkeypatch.setattr("builtins.print", _raise_runtime_error)

    with pytest.raises(RuntimeError):
        cli.main()

    log_content = (tmp_path / "logs" / "app.log").read_text(encoding="utf-8")
    assert "run.started" in log_content
    assert "run.finished" in log_content




def test_main_run_executes_extract_then_assess(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    output_dir = tmp_path / "data" / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    _write_extraction_xlsx(output_dir / "extraction_20260327T100000Z_olderrun1234.xlsx", [])

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "run", "--config", str(config_file)]
    )
    assert cli.main() == 0

    extract_files = list((tmp_path / "data" / "output").glob("extraction_*.xlsx"))
    assess_files = list((tmp_path / "data" / "output").glob("assessment_*.xlsx"))

    assert len(extract_files) >= 1
    assert len(assess_files) == 1
    assert all(extract_file.name != assess_files[0].name for extract_file in extract_files)

def test_main_creates_distinct_xlsx_artifacts_per_command(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    output_dir = tmp_path / "data" / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    _write_extraction_xlsx(output_dir / "extraction_20260327T100000Z_olderrun1234.xlsx", [])

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "extract", "--config", str(config_file)]
    )
    assert cli.main() == 0

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "assess", "--config", str(config_file)]
    )
    assert cli.main() == 0

    extract_files = list((tmp_path / "data" / "output").glob("extraction_*.xlsx"))
    assess_files = list((tmp_path / "data" / "output").glob("assessment_*.xlsx"))

    assert len(extract_files) >= 1
    assert len(assess_files) == 1
    assert all(extract_file.name != assess_files[0].name for extract_file in extract_files)


def test_main_assess_raises_clear_error_when_no_extraction_xlsx_exists(
    monkeypatch, tmp_path: Path
) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "assess", "--config", str(config_file)]
    )

    with pytest.raises(StorageIOError, match="Run 'extract' first"):
        cli.main()


def test_main_reads_and_updates_status_file(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    input_dir = tmp_path / "data" / "input"
    input_dir.mkdir(parents=True, exist_ok=True)
    (input_dir / "invoice_a.pdf").write_text("dummy", encoding="utf-8")

    class _FakeExtractionService:
        def extract_products_from_document(self, **kwargs: object) -> list[ExtractedProduct]:
            pdf_path = kwargs["pdf_path"]
            assert isinstance(pdf_path, Path)
            return [
                ExtractedProduct(
                    document_name=pdf_path.name,
                    product_name="Widget",
                    quantity="1",
                    supplier="ACME",
                    extraction_status="confirmed",
                    extraction_hint=None,
                )
            ]

    monkeypatch.setattr(
        cli, "_build_extraction_service", lambda **_kwargs: _FakeExtractionService()
    )
    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "extract", "--config", str(config_file)]
    )

    assert cli.main() == 0

    status_file = tmp_path / "data" / "state" / "processed_files.json"
    assert status_file.exists()

    payload = status_file.read_text(encoding="utf-8")
    assert '"file_name": "invoice_a.pdf"' in payload
    assert '"processed_at_utc":' in payload


def test_main_uses_status_fallback_strategy_from_config(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_FALLBACK_STATUS_CONFIG, encoding="utf-8")

    state_dir = tmp_path / "data" / "state"
    state_dir.mkdir(parents=True, exist_ok=True)
    (state_dir / "processed_files.json").write_text("{invalid json", encoding="utf-8")

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "extract", "--config", str(config_file)]
    )

    assert cli.main() == 0


def test_main_skips_already_processed_pdfs(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    input_dir = tmp_path / "data" / "input"
    input_dir.mkdir(parents=True, exist_ok=True)
    (input_dir / "invoice_old.pdf").write_text("dummy", encoding="utf-8")
    (input_dir / "invoice_new.pdf").write_text("dummy", encoding="utf-8")

    class _FakeExtractionService:
        def extract_products_from_document(self, **kwargs: object) -> list[ExtractedProduct]:
            pdf_path = kwargs["pdf_path"]
            assert isinstance(pdf_path, Path)
            return [
                ExtractedProduct(
                    document_name=pdf_path.name,
                    product_name="Widget",
                    quantity="1",
                    supplier="ACME",
                    extraction_status="confirmed",
                    extraction_hint=None,
                )
            ]

    monkeypatch.setattr(
        cli, "_build_extraction_service", lambda **_kwargs: _FakeExtractionService()
    )

    state_dir = tmp_path / "data" / "state"
    state_dir.mkdir(parents=True, exist_ok=True)
    (state_dir / "processed_files.json").write_text(
        json.dumps(
            {
                "processed_files": [
                    {
                        "file_name": "invoice_old.pdf",
                        "processed_at_utc": "2026-01-01T00:00:00Z",
                    }
                ]
            }
        ),
        encoding="utf-8",
    )

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "extract", "--config", str(config_file)]
    )

    assert cli.main() == 0

    status_file = tmp_path / "data" / "state" / "processed_files.json"
    payload = json.loads(status_file.read_text(encoding="utf-8"))
    file_names = [entry["file_name"] for entry in payload["processed_files"]]

    assert file_names == ["invoice_new.pdf", "invoice_old.pdf"]


def test_main_aborts_on_corrupt_status_by_default(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    state_dir = tmp_path / "data" / "state"
    state_dir.mkdir(parents=True, exist_ok=True)
    (state_dir / "processed_files.json").write_text("{invalid json", encoding="utf-8")

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "extract", "--config", str(config_file)]
    )

    with pytest.raises(StatusTrackingError, match="Could not read status file"):
        cli.main()


def test_main_extract_continues_after_single_document_failure(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    input_dir = tmp_path / "data" / "input"
    input_dir.mkdir(parents=True, exist_ok=True)
    (input_dir / "broken.pdf").write_text("dummy", encoding="utf-8")
    (input_dir / "partial.pdf").write_text("dummy", encoding="utf-8")

    class _FakeExtractionService:
        def extract_products_from_document(self, **kwargs: object) -> list[ExtractedProduct]:
            pdf_path = kwargs["pdf_path"]
            assert isinstance(pdf_path, Path)
            if pdf_path.name == "broken.pdf":
                raise PdfProcessingError("broken document")
            return [
                ExtractedProduct(
                    document_name=pdf_path.name,
                    product_name="Widget",
                    quantity="UNKNOWN",
                    supplier="UNKNOWN",
                    extraction_status="uncertain",
                    extraction_hint="Missing required fields: quantity, supplier",
                )
            ]

    monkeypatch.setattr(
        cli, "_build_extraction_service", lambda **_kwargs: _FakeExtractionService()
    )
    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "extract", "--config", str(config_file)]
    )

    assert cli.main() == 0

    output_files = list((tmp_path / "data" / "output").glob("extraction_*.xlsx"))
    assert len(output_files) == 1

    rows = _read_xlsx_rows(output_files[0])
    text_dump = "\n".join("|".join(str(cell or "") for cell in row) for row in rows)
    assert "broken.pdf" in text_dump
    assert "partial.pdf" in text_dump
    assert "uncertain" in text_dump
    assert "Document processing failed: PdfProcessingError" in text_dump
    assert "Missing required fields: quantity, supplier" in text_dump

    status_file = tmp_path / "data" / "state" / "processed_files.json"
    payload = json.loads(status_file.read_text(encoding="utf-8"))
    assert [entry["file_name"] for entry in payload["processed_files"]] == ["partial.pdf"]


def test_main_extract_persists_status_immediately_after_success(
    monkeypatch, tmp_path: Path
) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    input_dir = tmp_path / "data" / "input"
    input_dir.mkdir(parents=True, exist_ok=True)
    (input_dir / "invoice_a.pdf").write_text("dummy", encoding="utf-8")
    (input_dir / "invoice_b.pdf").write_text("dummy", encoding="utf-8")

    status_file = tmp_path / "data" / "state" / "processed_files.json"

    class _FakeExtractionService:
        def extract_products_from_document(self, **kwargs: object) -> list[ExtractedProduct]:
            pdf_path = kwargs["pdf_path"]
            assert isinstance(pdf_path, Path)

            if pdf_path.name == "invoice_b.pdf":
                status_payload = json.loads(status_file.read_text(encoding="utf-8"))
                assert [entry["file_name"] for entry in status_payload["processed_files"]] == [
                    "invoice_a.pdf"
                ]

            return [
                ExtractedProduct(
                    document_name=pdf_path.name,
                    product_name="Widget",
                    quantity="1",
                    supplier="ACME",
                    extraction_status="confirmed",
                    extraction_hint=None,
                )
            ]

    monkeypatch.setattr(
        cli, "_build_extraction_service", lambda **_kwargs: _FakeExtractionService()
    )
    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "extract", "--config", str(config_file)]
    )

    assert cli.main() == 0


def test_main_assess_marks_unprocessed_pdfs_as_processed(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    input_dir = tmp_path / "data" / "input"
    input_dir.mkdir(parents=True, exist_ok=True)
    (input_dir / "invoice_1.pdf").write_text("dummy", encoding="utf-8")

    output_dir = tmp_path / "data" / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    _write_extraction_xlsx(output_dir / "extraction_20260327T100000Z_run1234567890.xlsx", [])

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "assess", "--config", str(config_file)]
    )
    assert cli.main() == 0

    status_file = tmp_path / "data" / "state" / "processed_files.json"
    payload = json.loads(status_file.read_text(encoding="utf-8"))
    assert [entry["file_name"] for entry in payload["processed_files"]] == ["invoice_1.pdf"]


def test_invoke_extraction_llm_placeholder_raises_clear_runtime_error() -> None:
    with pytest.raises(RuntimeError, match="LLM provider invocation is not configured"):
        cli._invoke_extraction_llm_placeholder("prompt")


def test_invoke_assessment_llm_placeholder_raises_clear_runtime_error() -> None:
    with pytest.raises(RuntimeError, match="LLM provider invocation is not configured"):
        cli._invoke_assessment_llm_placeholder("prompt")


def test_main_assess_keeps_skipped_products_in_output_xlsx(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    output_dir = tmp_path / "data" / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    _write_extraction_xlsx(
        output_dir / "extraction_20260327T100000Z_run1234567890.xlsx",
        [["run123", "invoice_1.pdf", "UNKNOWN", "10", "ACME", "", "", "uncertain", ""]],
    )

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "assess", "--config", str(config_file)]
    )
    assert cli.main() == 0

    assessment_files = list(output_dir.glob("assessment_*.xlsx"))
    assert len(assessment_files) == 1
    rows = _read_xlsx_rows(assessment_files[0])
    text_dump = "\n".join("|".join(str(cell or "") for cell in row) for row in rows)
    assert "skipped" in text_dump
    assert "UNCONFIRMED_EXTRACTION" in text_dump


def test_main_assess_prints_operational_console_summary(
    monkeypatch,
    tmp_path: Path,
    capsys,
) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    output_dir = tmp_path / "data" / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    _write_extraction_xlsx(
        output_dir / "extraction_20260327T100000Z_run1234567890.xlsx",
        [["run123", "invoice_1.pdf", "UNKNOWN", "10", "ACME", "", "", "uncertain", ""]],
    )

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "assess", "--config", str(config_file)]
    )
    assert cli.main() == 0

    stdout = capsys.readouterr().out
    assert "Assessment results" in stdout
    assert "Produktname: UNKNOWN" in stdout
    assert "Lieferant: ACME" in stdout
    assert "Risikostufe: -" in stdout
    assert "Preisänderung: -" in stdout
    assert "Status: skipped (UNCONFIRMED_EXTRACTION)" in stdout


def test_main_assess_writes_new_xlsx_for_each_run(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    output_dir = tmp_path / "data" / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    _write_extraction_xlsx(
        output_dir / "extraction_20260327T100000Z_run1234567890.xlsx",
        [["run123", "invoice_1.pdf", "UNKNOWN", "10", "ACME", "", "", "uncertain", ""]],
    )

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "assess", "--config", str(config_file)]
    )
    assert cli.main() == 0

    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "assess", "--config", str(config_file)]
    )
    assert cli.main() == 0

    assessment_files = list(output_dir.glob("assessment_*.xlsx"))
    assert len(assessment_files) == 2
    assert assessment_files[0].name != assessment_files[1].name


def test_print_assessment_console_results_handles_empty_results(capsys) -> None:
    cli._print_assessment_console_results(results=[])

    output = capsys.readouterr().out
    assert "Assessment results" in output
    assert "- keine Produkte vorhanden" in output


def test_main_extract_aborts_on_llm_configuration_errors(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.chdir(tmp_path)

    config_file = tmp_path / "config.yaml"
    config_file.write_text(_MINIMAL_CONFIG, encoding="utf-8")

    input_dir = tmp_path / "data" / "input"
    input_dir.mkdir(parents=True, exist_ok=True)
    (input_dir / "invoice_a.pdf").write_text("dummy", encoding="utf-8")

    class _MisconfiguredExtractionService:
        def extract_products_from_document(self, **_kwargs: object) -> list[ExtractedProduct]:
            raise LlmConfigurationError("OPENAI_API_KEY is required")

    monkeypatch.setattr(
        cli, "_build_extraction_service", lambda **_kwargs: _MisconfiguredExtractionService()
    )
    monkeypatch.setattr(
        "sys.argv", ["supply-chain-checker", "extract", "--config", str(config_file)]
    )

    with pytest.raises(LlmConfigurationError, match="OPENAI_API_KEY"):
        cli.main()
