"""Tests for XLSX artefact creation."""

from __future__ import annotations

from datetime import UTC, datetime

import pytest
from openpyxl import Workbook, load_workbook

from supply_chain_checker.models import ExtractedProduct
from supply_chain_checker.run_context import RunContext
from supply_chain_checker.services.assessment_service import ProductAssessmentResult
from supply_chain_checker.services.csv_service import (
    StorageIOError,
    build_run_csv_filename,
    create_run_csv_artifact,
    read_extraction_products_csv,
    select_latest_extraction_csv,
    write_assessment_results_csv,
)


def _rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    return list(wb.active.iter_rows(values_only=True))


def test_build_run_csv_filename_uses_domain_prefix_and_run_metadata() -> None:
    context = RunContext(
        run_id="abc123def456",
        started_at_utc=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
    )

    extraction_name = build_run_csv_filename(command="extract", run_context=context)
    assessment_name = build_run_csv_filename(command="assess", run_context=context)

    assert extraction_name == "extraction_20260327T120000Z_abc123def456.xlsx"
    assert assessment_name == "assessment_20260327T120000Z_abc123def456.xlsx"


def test_create_run_csv_artifact_contains_run_id_for_log_correlation(tmp_path) -> None:
    context = RunContext(
        run_id="runid1234567",
        started_at_utc=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
    )

    artifact_path = create_run_csv_artifact(
        output_dir=tmp_path,
        command="extract",
        run_context=context,
    )

    rows = _rows(artifact_path)
    header = rows[0]
    values = rows[1]

    assert header[0] == "run_id"
    assert values[0] == "runid1234567"
    assert values[2] == "extract"


def test_select_latest_extraction_csv_uses_timestamp_in_file_name(tmp_path) -> None:
    older = tmp_path / "extraction_20260327T100000Z_runolder1234.xlsx"
    newer = tmp_path / "extraction_20260327T120000Z_runnewer1234.xlsx"
    ignored = tmp_path / "extraction_invalid.xlsx"

    older.write_bytes(b"old")
    newer.write_bytes(b"new")
    ignored.write_bytes(b"ignored")

    selected = select_latest_extraction_csv(output_dir=tmp_path)

    assert selected == newer


def test_select_latest_extraction_csv_raises_when_no_valid_file_exists(tmp_path) -> None:
    (tmp_path / "assessment_20260327T120000Z_run1234.xlsx").write_bytes(b"")

    with pytest.raises(StorageIOError, match="Run 'extract' first"):
        select_latest_extraction_csv(output_dir=tmp_path)


def test_read_extraction_products_csv_loads_domain_rows(tmp_path) -> None:
    xlsx_path = tmp_path / "extraction_20260327T120000Z_run123.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "run_id",
            "document_name",
            "product_name",
            "quantity",
            "supplier",
            "manufacturer",
            "article_number",
            "extraction_status",
            "extraction_hint",
        ]
    )
    ws.append(["run123", "invoice.pdf", "Bolt", "5", "ACME", "SupplierCo", "ART-1", "confirmed", ""])
    wb.save(xlsx_path)

    products = read_extraction_products_csv(csv_path=xlsx_path)

    assert len(products) == 1
    assert products[0].product_name == "Bolt"
    assert products[0].extraction_status == "confirmed"


def test_write_assessment_results_csv_keeps_skipped_products_and_skip_reason(tmp_path) -> None:
    context = RunContext(
        run_id="runid1234567",
        started_at_utc=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
    )
    result = ProductAssessmentResult(
        product=ExtractedProduct(
            document_name="invoice.pdf",
            product_name="UNKNOWN",
            quantity="5",
            supplier="ACME",
            extraction_status="uncertain",
        ),
        assessment_status="skipped",
        raw_response=None,
        skip_reason="UNCONFIRMED_EXTRACTION",
    )

    xlsx_path = write_assessment_results_csv(
        output_dir=tmp_path,
        run_context=context,
        results=[result],
    )
    rows = _rows(xlsx_path)

    assert rows[1][9] == "skipped"
    assert rows[1][10] == "UNCONFIRMED_EXTRACTION"


def test_read_extraction_products_csv_wraps_io_errors(tmp_path, monkeypatch) -> None:
    xlsx_path = tmp_path / "extraction_20260327T120000Z_run123.xlsx"
    xlsx_path.write_bytes(b"not-a-workbook")

    def _raise_oserror(*_args, **_kwargs):
        raise OSError("boom")

    monkeypatch.setattr("supply_chain_checker.services.csv_service.load_workbook", _raise_oserror)

    with pytest.raises(StorageIOError, match="Could not read extraction XLSX"):
        read_extraction_products_csv(csv_path=xlsx_path)


def test_write_assessment_results_csv_wraps_io_errors(tmp_path, monkeypatch) -> None:
    context = RunContext(
        run_id="runid1234567",
        started_at_utc=datetime(2026, 3, 27, 12, 0, tzinfo=UTC),
    )

    result = ProductAssessmentResult(
        product=ExtractedProduct(
            document_name="invoice.pdf",
            product_name="Bolt",
            quantity="5",
            supplier="ACME",
        ),
        assessment_status="assessed",
        raw_response='{"ok": true}',
    )

    def _raise_save(self, _path):
        raise OSError("write boom")

    monkeypatch.setattr("openpyxl.workbook.workbook.Workbook.save", _raise_save)

    with pytest.raises(StorageIOError, match="Could not write assessment XLSX"):
        write_assessment_results_csv(output_dir=tmp_path, run_context=context, results=[result])


def test_select_latest_extraction_csv_breaks_ties_by_filename(tmp_path) -> None:
    first = tmp_path / "extraction_20260327T120000Z_arun.xlsx"
    second = tmp_path / "extraction_20260327T120000Z_zrun.xlsx"
    first.write_bytes(b"a")
    second.write_bytes(b"z")

    selected = select_latest_extraction_csv(output_dir=tmp_path)

    assert selected == second


def test_select_latest_extraction_csv_ignores_files_with_invalid_pattern(tmp_path) -> None:
    (tmp_path / "extraction_2026-03-27T120000Z_bad.xlsx").write_bytes(b"bad")
    (tmp_path / "extraction_20260327T120000_run.xlsx").write_bytes(b"bad")
    (tmp_path / "extraction_20260327T120000Z.xlsx").write_bytes(b"bad")

    with pytest.raises(StorageIOError, match="Run 'extract' first"):
        select_latest_extraction_csv(output_dir=tmp_path)
